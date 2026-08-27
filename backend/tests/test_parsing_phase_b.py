import io

import docx
import pymupdf
import pytest
from docx.enum.text import WD_BREAK
from PIL import Image, ImageDraw

from app.parsing.docx import DocxParser
from app.parsing.errors import DocumentParseError, UnsupportedDocumentError
from app.parsing.factory import get_parser
from app.parsing.image_ocr import ImageOcrParser
from app.parsing.pdf import PdfParser


# ----------------------------------------------------------------------
# Fixtures / helpers
# ----------------------------------------------------------------------

def build_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Page one text.")

    break_paragraph = document.add_paragraph()
    break_paragraph.add_run().add_break(WD_BREAK.PAGE)

    document.add_paragraph("Page two text.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Vendor"
    table.cell(0, 1).text = "Acme Supplies"
    table.cell(1, 0).text = "Contract Value"
    table.cell(1, 1).text = "100000"

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def build_docx_bytes_no_page_break() -> bytes:
    document = docx.Document()
    document.add_paragraph("Only paragraph, no explicit page break.")

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def render_text_to_png(text: str) -> bytes:
    image = Image.new("RGB", (900, 200), color="white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 80), text, fill="black")

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    return buffer.getvalue()


def build_image_only_pdf_bytes(text: str) -> bytes:
    """A PDF page containing only a rendered image -- no text layer at
    all -- to force the OCR fallback path in PdfParser."""

    image_bytes = render_text_to_png(text)

    document = pymupdf.open()
    page = document.new_page(width=900, height=200)
    page.insert_image(pymupdf.Rect(0, 0, 900, 200), stream=image_bytes)

    content = document.tobytes()
    document.close()

    return content


# ----------------------------------------------------------------------
# DocxParser
# ----------------------------------------------------------------------

def test_docx_parser_supports_docx():
    parser = DocxParser()

    assert parser.supports(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "contract.docx",
    )


def test_docx_parser_supports_extension_fallback():
    parser = DocxParser()

    assert parser.supports("application/octet-stream", "contract.docx")


def test_docx_parser_rejects_non_docx():
    parser = DocxParser()

    assert not parser.supports("application/pdf", "contract.pdf")


def test_docx_parser_extracts_pages_split_on_page_breaks():
    parser = DocxParser()

    result = parser.parse(
        content=build_docx_bytes(),
        filename="contract.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )

    body_pages = [
        p for p in result.pages if p.metadata["source_type"] == "docx"
    ]

    assert len(body_pages) == 2
    assert body_pages[0].index == 0
    assert "Page one text." in body_pages[0].text
    assert body_pages[1].index == 1
    assert "Page two text." in body_pages[1].text

    assert result.metadata["parser"] == "docx"


def test_docx_parser_extracts_tables_as_separate_pages():
    parser = DocxParser()

    result = parser.parse(
        content=build_docx_bytes(),
        filename="contract.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )

    table_pages = [
        p for p in result.pages if p.metadata["source_type"] == "docx_table"
    ]

    assert len(table_pages) == 1
    assert "Vendor" in table_pages[0].text
    assert "Acme Supplies" in table_pages[0].text
    assert "Contract Value" in table_pages[0].text
    assert "100000" in table_pages[0].text


def test_docx_parser_single_page_when_no_page_breaks():
    parser = DocxParser()

    result = parser.parse(
        content=build_docx_bytes_no_page_break(),
        filename="simple.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )

    body_pages = [
        p for p in result.pages if p.metadata["source_type"] == "docx"
    ]

    assert len(body_pages) == 1
    assert body_pages[0].index == 0
    assert "Only paragraph" in body_pages[0].text


def test_docx_parser_rejects_invalid_docx():
    parser = DocxParser()

    with pytest.raises(DocumentParseError):
        parser.parse(
            content=b"this is not a docx file",
            filename="broken.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document"
            ),
        )


# ----------------------------------------------------------------------
# PdfParser OCR fallback
# ----------------------------------------------------------------------

def test_pdf_parser_ocr_fallback_on_scanned_page():
    parser = PdfParser()

    content = build_image_only_pdf_bytes("Vendor Acme Supplies")

    result = parser.parse(
        content=content,
        filename="scan.pdf",
        content_type="application/pdf",
    )

    assert len(result.pages) == 1
    assert result.pages[0].metadata["source_type"] == "pdf_ocr"
    assert "Acme" in result.pages[0].text
    assert result.metadata["ocr_page_count"] == 1


def test_pdf_parser_does_not_ocr_pages_with_real_text():
    parser = PdfParser()

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Real embedded text, not a scan.")
    content = document.tobytes()
    document.close()

    result = parser.parse(
        content=content,
        filename="real.pdf",
        content_type="application/pdf",
    )

    assert result.pages[0].metadata["source_type"] == "pdf"
    assert result.metadata["ocr_page_count"] == 0


# ----------------------------------------------------------------------
# ImageOcrParser
# ----------------------------------------------------------------------

def test_image_ocr_parser_supports_common_image_types():
    parser = ImageOcrParser()

    assert parser.supports("image/png", "scan.png")
    assert parser.supports("image/jpeg", "scan.jpg")
    assert parser.supports("application/octet-stream", "scan.tiff")


def test_image_ocr_parser_rejects_non_image():
    parser = ImageOcrParser()

    assert not parser.supports("application/pdf", "doc.pdf")


def test_image_ocr_parser_extracts_text():
    parser = ImageOcrParser()

    content = render_text_to_png("Contract Value 100000")

    result = parser.parse(
        content=content,
        filename="scan.png",
        content_type="image/png",
    )

    assert "Contract Value" in result.text
    assert len(result.pages) == 1
    assert result.pages[0].metadata["source_type"] == "image_ocr"
    assert result.metadata["parser"] == "image_ocr"


def test_image_ocr_parser_rejects_invalid_image():
    parser = ImageOcrParser()

    with pytest.raises(DocumentParseError):
        parser.parse(
            content=b"not an image",
            filename="broken.png",
            content_type="image/png",
        )


# ----------------------------------------------------------------------
# Factory wiring
# ----------------------------------------------------------------------

def test_factory_returns_docx_parser():
    parser = get_parser(
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
        filename="contract.docx",
    )

    assert isinstance(parser, DocxParser)


def test_factory_returns_image_ocr_parser():
    parser = get_parser(content_type="image/png", filename="scan.png")

    assert isinstance(parser, ImageOcrParser)


def test_factory_still_rejects_truly_unknown_formats():
    with pytest.raises(UnsupportedDocumentError):
        get_parser(content_type="application/x-unknown", filename="mystery.xyz")
