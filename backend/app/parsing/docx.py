import io

import docx

from app.parsing.base import DocumentParser
from app.parsing.errors import DocumentParseError
from app.parsing.models import ParsedDocument, ParsedPage


class DocxParser(DocumentParser):
    """Parser for Word (.docx) documents.

    DOCX has no fixed "page" concept the way PDF does -- page breaks
    are a rendering-time detail, not stored layout. This parser
    recovers approximate page boundaries from explicit page-break runs
    (`<w:br w:type="page"/>`), which is what a human inserting
    Ctrl+Enter/"Insert Page Break" produces and is the common case for
    contracts and similar formal documents. A document with no
    explicit page breaks is treated as a single page (index 0) --
    accurate page-aware evidence still degrades gracefully to
    document-aware evidence in that case, it never silently invents a
    page number.

    Tables are extracted separately and appended as their own
    "page"(s) after the body text, since contract/invoice tables
    (vendor/amount/date grids) are exactly the kind of content this
    domain's rule and extraction logic needs and paragraph iteration
    alone skips them entirely.
    """

    SUPPORTED_CONTENT_TYPES = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }

    SUPPORTED_EXTENSIONS = {
        ".docx",
    }

    def supports(self, content_type: str, filename: str) -> bool:
        normalized_content_type = content_type.lower().split(";")[0].strip()
        normalized_filename = filename.lower()

        return (
            normalized_content_type in self.SUPPORTED_CONTENT_TYPES
            or normalized_filename.endswith(tuple(self.SUPPORTED_EXTENSIONS))
        )

    def parse(
        self,
        *,
        content: bytes,
        filename: str,
        content_type: str,
    ) -> ParsedDocument:
        try:
            document = docx.Document(io.BytesIO(content))
        except Exception as exc:
            raise DocumentParseError(
                f"Unable to open DOCX document: {filename}"
            ) from exc

        try:
            pages = self._extract_pages(document)
            pages += self._extract_table_pages(document, start_index=len(pages))

            if not pages:
                pages = [
                    ParsedPage(
                        index=0,
                        text="",
                        metadata={"source_type": "docx"},
                    )
                ]

            full_text = "\n".join(page.text for page in pages)

            return ParsedDocument(
                text=full_text,
                filename=filename,
                content_type=content_type,
                pages=tuple(pages),
                metadata={
                    "parser": "docx",
                    "page_count": len(pages),
                },
            )
        except Exception as exc:
            raise DocumentParseError(
                f"Unable to parse DOCX document: {filename}"
            ) from exc

    def _extract_pages(self, document) -> list[ParsedPage]:
        pages: list[ParsedPage] = []
        current_lines: list[str] = []
        page_index = 0

        for paragraph in document.paragraphs:
            has_page_break = bool(
                paragraph._element.xpath('.//w:br[@w:type="page"]')
            )

            if paragraph.text.strip():
                current_lines.append(paragraph.text)

            if has_page_break:
                pages.append(
                    ParsedPage(
                        index=page_index,
                        text="\n".join(current_lines),
                        metadata={
                            "source_type": "docx",
                            "page_number": page_index + 1,
                        },
                    )
                )
                current_lines = []
                page_index += 1

        # Final (or only, if no page breaks) page.
        if current_lines or page_index == 0:
            pages.append(
                ParsedPage(
                    index=page_index,
                    text="\n".join(current_lines),
                    metadata={
                        "source_type": "docx",
                        "page_number": page_index + 1,
                    },
                )
            )

        return pages

    def _extract_table_pages(
        self,
        document,
        *,
        start_index: int,
    ) -> list[ParsedPage]:
        pages: list[ParsedPage] = []

        for table_number, table in enumerate(document.tables):
            rows_text = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]

            if not any(rows_text):
                continue

            pages.append(
                ParsedPage(
                    index=start_index + table_number,
                    text="\n".join(rows_text),
                    metadata={
                        "source_type": "docx_table",
                        "table_number": table_number + 1,
                    },
                )
            )

        return pages
